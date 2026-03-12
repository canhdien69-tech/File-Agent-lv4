# ==============================================================================
#  SUNNY AI v5.0 — tools/reader.py
#  File reader với RAM guard + sandbox check.
# ==============================================================================
from core.sandbox import Sandbox

try:
    import PyPDF2
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


MAX_PDF_PAGES  = 8
MAX_DOCX_PARAS = 80
MAX_XLSX_ROWS  = 40
MAX_CHARS      = 6000


def _clean(text: str) -> str:
    return "\n".join(ln.strip() for ln in text.split("\n") if ln.strip())


def read_file(path: str) -> str:
    ok, result = Sandbox.check_path(path)
    if not ok:
        return f"[Blocked]: {result}"

    resolved = result   # Sandbox trả về resolved path khi ok=True
    ext      = __import__("os").path.splitext(resolved)[1].lower()
    content  = ""

    try:
        if ext == ".pdf":
            if not HAS_PDF:
                return "Missing PyPDF2. Run: pip install PyPDF2"
            reader = PyPDF2.PdfReader(resolved)
            for page in reader.pages[:MAX_PDF_PAGES]:
                content += (page.extract_text() or "") + "\n"

        elif ext == ".docx":
            if not HAS_DOCX:
                return "Missing python-docx. Run: pip install python-docx"
            for para in DocxDocument(resolved).paragraphs[:MAX_DOCX_PARAS]:
                content += para.text + "\n"

        elif ext in (".xlsx", ".xls"):
            if HAS_PANDAS:
                content = pd.read_excel(resolved, nrows=MAX_XLSX_ROWS).to_string(index=False)
            elif HAS_OPENPYXL:
                wb = openpyxl.load_workbook(resolved, data_only=True)
                for sheet in wb:
                    content += f"\n[Sheet: {sheet.title}]\n"
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= MAX_XLSX_ROWS:
                            break
                        row_d = [str(c) for c in row if c is not None]
                        if row_d:
                            content += " | ".join(row_d) + "\n"
            else:
                return "Missing pandas/openpyxl."

        elif ext in (".txt", ".csv", ".md"):
            with open(resolved, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

        return _clean(content)[:MAX_CHARS]

    except Exception as e:
        return f"Read error: {e}"
